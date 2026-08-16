"""File tools for the Controlled Agent Platform."""

from __future__ import annotations

import asyncio
from pathlib import Path
from typing import Any

from packages.policy.unified_registry import tool_registry

from .base import ExecutionContext, ToolBase, ToolResult

MAX_FILE_SIZE_MB_DEFAULT = 50
_MAX_CONTENT_CHARS = MAX_FILE_SIZE_MB_DEFAULT * 1024 * 1024


def _check_path_within_base(path: Path, base: Path) -> tuple[bool, str]:
    """Check that *path* resolves inside *base*. Returns (safe, error_msg).

    Also blocks symbolic links to prevent workspace escape via symlink chains.
    """
    try:
        if path.is_symlink() or any(p.is_symlink() for p in path.parents if p != Path("/")):
            return False, "Symbolic links not allowed"
        resolved = path.resolve()
        base_resolved = base.resolve()
    except OSError:
        return False, "Cannot resolve path"
    if not resolved.is_relative_to(base_resolved):
        return False, "Path escapes workspace"
    return True, ""


def _check_file_size(path: Path, max_mb: int) -> tuple[bool, str]:
    """Check that *path* does not exceed *max_mb*. Returns (ok, error_msg)."""
    size_bytes = path.stat().st_size
    limit = max_mb * 1024 * 1024
    if size_bytes > limit:
        return False, f"File size {size_bytes} bytes exceeds limit of {max_mb} MB"
    return True, ""


@tool_registry.register(
    category="file",
    risk_level="low",
    emoji="📝",
    params_schema={
        "type": "object",
        "required": ["output_path", "title", "paragraphs"],
        "properties": {
            "output_path": {"type": "string", "description": "Relative path under workspace/outputs"},
            "title": {"type": "string"},
            "paragraphs": {"type": "array", "items": {"type": "string"}},
            "tables": {"type": "array", "items": {"type": "object"}},
        },
        "additionalProperties": False,
    },
)
class FileWriteDocx(ToolBase):
    """Write content to a Word (.docx) document."""

    name = "file.write_docx"
    description = "Write content to a Word document"

    async def execute(self, args: dict[str, Any], context: ExecutionContext) -> ToolResult:
        output_path = args.get("output_path", "output.docx")
        title = args.get("title", "Document")
        paragraphs: list[str] = args.get("paragraphs", [])
        tables: list[dict[str, Any]] = args.get("tables", [])

        # Reject workspace escapes before probing optional dependencies.
        base = Path(context.outputs_path)
        full_path = base / output_path
        safe, err = _check_path_within_base(full_path, base)
        if not safe:
            return ToolResult(success=False, error=err)

        try:
            from docx import Document  # type: ignore[import-untyped]
        except ImportError:
            return ToolResult(
                success=False,
                error="python-docx not installed. Install with: pip install python-docx",
            )

        # Build document
        doc = Document()
        doc.add_heading(title, level=1)
        for p in paragraphs:
            doc.add_paragraph(p)
        for t in tables:
            rows: list[list[Any]] = t.get("rows", [])
            if rows:
                num_cols = len(rows[0])
                table = doc.add_table(rows=len(rows), cols=num_cols)
                for i, row in enumerate(rows):
                    for j, cell in enumerate(row):
                        table.rows[i].cells[j].text = str(cell)

        # Save
        full_path.parent.mkdir(parents=True, exist_ok=True)
        await asyncio.to_thread(doc.save, str(full_path))

        # File size check
        max_mb = context.max_file_size_mb or MAX_FILE_SIZE_MB_DEFAULT
        ok, size_err = _check_file_size(full_path, max_mb)
        if not ok:
            full_path.unlink(missing_ok=True)
            return ToolResult(success=False, error=size_err)

        return ToolResult(
            success=True,
            output=str(full_path),
            artifacts=[str(full_path)],
        )


@tool_registry.register(
    category="file",
    risk_level="low",
    emoji="📋",
    params_schema={
        "type": "object",
        "required": ["output_path", "content"],
        "properties": {
            "output_path": {"type": "string"},
            "content": {"type": "string"},
        },
        "additionalProperties": False,
    },
)
class FileWriteMarkdown(ToolBase):
    """Write content to a Markdown (.md) file."""

    name = "file.write_markdown"
    description = "Write content to a Markdown file"

    async def execute(self, args: dict[str, Any], context: ExecutionContext) -> ToolResult:
        output_path = args.get("output_path", "output.md")
        content: str = args.get("content", "")

        max_mb = context.max_file_size_mb or MAX_FILE_SIZE_MB_DEFAULT
        max_chars = max_mb * 1024 * 1024
        if len(content) > max_chars:
            return ToolResult(success=False, error=f"Content length {len(content)} chars exceeds limit of {max_mb} MB")

        # Workspace safety check
        base = Path(context.outputs_path)
        full_path = base / output_path
        safe, err = _check_path_within_base(full_path, base)
        if not safe:
            return ToolResult(success=False, error=err)

        # Write
        full_path.parent.mkdir(parents=True, exist_ok=True)
        await asyncio.to_thread(full_path.write_text, content, encoding="utf-8")

        # File size check
        max_mb = context.max_file_size_mb or MAX_FILE_SIZE_MB_DEFAULT
        ok, size_err = _check_file_size(full_path, max_mb)
        if not ok:
            full_path.unlink(missing_ok=True)
            return ToolResult(success=False, error=size_err)

        return ToolResult(
            success=True,
            output=str(full_path),
            artifacts=[str(full_path)],
        )


@tool_registry.register(
    category="file",
    risk_level="low",
    emoji="📖",
    params_schema={
        "type": "object",
        "required": ["path"],
        "properties": {
            "path": {"type": "string", "description": "Relative path under workspace"},
        },
        "additionalProperties": False,
    },
)
class FileRead(ToolBase):
    """Read text content from a file in the workspace."""

    name = "file.read"
    description = "Read a file content"

    async def execute(self, args: dict[str, Any], context: ExecutionContext) -> ToolResult:
        rel_path: str = args.get("path", "")
        if not rel_path:
            return ToolResult(success=False, error="path is required")

        # Workspace safety check
        base = Path(context.workspace_root)
        full_path = base / rel_path
        safe, err = _check_path_within_base(full_path, base)
        if not safe:
            return ToolResult(success=False, error=err)

        if not full_path.is_file():
            return ToolResult(success=False, error="File not found")

        max_mb = context.max_file_size_mb or MAX_FILE_SIZE_MB_DEFAULT
        ok, size_err = _check_file_size(full_path, max_mb)
        if not ok:
            return ToolResult(success=False, error=size_err)

        try:
            content = await asyncio.to_thread(full_path.read_text, encoding="utf-8")
        except OSError as exc:
            return ToolResult(success=False, error=f"Failed to read file: {exc}")

        return ToolResult(success=True, output=content)


@tool_registry.register(
    category="file",
    risk_level="low",
    emoji="📁",
    params_schema={
        "type": "object",
        "properties": {
            "path": {"type": "string", "default": "."},
            "pattern": {"type": "string", "default": "*"},
        },
        "additionalProperties": False,
    },
)
class FileList(ToolBase):
    """List files in a directory under the workspace."""

    name = "file.list"
    description = "List files in a directory"

    async def execute(self, args: dict[str, Any], context: ExecutionContext) -> ToolResult:
        rel_path: str = args.get("path") or "."
        pattern: str = args.get("pattern") or "*"

        if ".." in pattern or pattern.startswith("/") or pattern.startswith("\\"):
            return ToolResult(success=False, error="Invalid pattern")
        if "**" in pattern:
            return ToolResult(success=False, error="Recursive patterns not allowed")

        # Workspace safety check
        base = Path(context.workspace_root)
        full_path = (base / rel_path).resolve()
        safe, err = _check_path_within_base(full_path, base)
        if not safe:
            return ToolResult(success=False, error=err)

        if not full_path.is_dir():
            return ToolResult(success=False, error="Directory not found")

        try:
            base_resolved = base.resolve()

            def _list_files():
                return sorted(
                    str(p.resolve().relative_to(base_resolved))
                    for p in full_path.glob(pattern)
                    if p.is_file() and p.resolve().is_relative_to(base_resolved)
                )

            entries = await asyncio.to_thread(_list_files)
        except OSError as exc:
            return ToolResult(success=False, error=f"Failed to list files: {exc}")

        return ToolResult(success=True, output=entries)
